from datetime import datetime
import io
import re
from flask import Flask, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pdfplumber

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024

# Configurações de colunas e regex
COLUNAS = [
    'Conhec.',
    'Emissão',
    'Origem',
    'Destino',
    'Motorista',
    'Mercadoria',
    'Quant',
    'Veículo',
    'No. NF',
    'Vl. Seguro',
    'Desp. Extra',
    'Vl. ICMS',
    'Vl. Ped.',
    'Peso Saída',
    'Peso Cheg.',
    'PPT Emp.',
    'PPT Mot.',
    'Adt. Mot.',
    'Fr. Emp',
    'Fr. Mot',
    'Base Calc.',
    'Fr. Fiscal',
    'Margem',
    '% Margem',
]

DATE_RE = re.compile(r'^\d{2}/\d{2}/\d{4}$')
INT_RE = re.compile(r'^\d+$')
LIMITES_DIREITA = [420, 464, 508, 551, 595, 639, 695, 742]
MEIOS = [
    (LIMITES_DIREITA[i] + LIMITES_DIREITA[i + 1]) / 2
    for i in range(len(LIMITES_DIREITA) - 1)
]


def normalizar_texto(texto):
  return re.sub(r'\s+', ' ', texto or '').strip()


def agrupar_linhas(page):
  palavras = page.extract_words(
      x_tolerance=1, y_tolerance=2, keep_blank_chars=False
  )
  grupos = []
  for palavra in sorted(palavras, key=lambda item: (item['top'], item['x0'])):
    if not grupos or abs(palavra['top'] - grupos[-1][0]) > 1.5:
      grupos.append([palavra['top'], []])
    grupos[-1][1].append(palavra)
  return [(y, sorted(words, key=lambda item: item['x0'])) for y, words in grupos]


def texto_linha(words):
  return normalizar_texto(' '.join(word['text'] for word in words))


def texto_por_x(words, x_min, x_max=None):
  valores = [
      word['text']
      for word in words
      if word['x0'] >= x_min and (x_max is None or word['x0'] < x_max)
  ]
  return normalizar_texto(' '.join(valores))


def valor_numerico(texto):
  if not texto or ' ' in texto.strip():
    return None
  texto = texto.strip()
  try:
    if ',' in texto:
      return float(texto.replace('.', '').replace(',', '.'))
    if re.fullmatch(r'-?\d+\.\d+', texto):
      return float(texto)
    if re.fullmatch(r'-?\d+', texto):
      return int(texto)
  except ValueError:
    pass
  return texto


def campo_numerico(words, indice):
  limite_inferior = 390 if indice == 0 else MEIOS[indice - 1]
  limite_superior = MEIOS[indice] if indice < len(MEIOS) else 750
  candidatos = [
      word
      for word in words
      if word['x1'] > limite_inferior and word['x1'] <= limite_superior
  ]
  if not candidatos:
    return None
  palavra = max(candidatos, key=lambda item: item['x1'])
  return valor_numerico(palavra['text'])


def eh_linha_data(words):
  x22 = [word['text'] for word in words if word['x0'] < 50]
  return (
      bool(x22)
      and bool(DATE_RE.fullmatch(x22[0]))
      and bool(texto_por_x(words, 60, 170))
      and bool(texto_por_x(words, 170, 300))
  )


def eh_possivel_primeira_linha(words):
  texto = texto_linha(words)
  ignorar = (
      'Proprietário:',
      'Totais - Viagens:',
      'RIO BONITO TRANSPORTES LTDA',
      'Relatório',
      'Conhec.',
      'Endereço:',
  )
  if texto.startswith(ignorar):
    return False
  return bool(
      texto_por_x(words, 60, 170)
      or texto_por_x(words, 170, 330)
      or texto_por_x(words, 330, 382)
  )


def processar_pdf_stream(pdf_file_stream):
  registros = []
  proprietario_atual = None
  pending_first_line = None

  with pdfplumber.open(pdf_file_stream) as pdf:
    for page in pdf.pages:
      grupos = agrupar_linhas(page)
      for indice, (_, words) in enumerate(grupos):
        texto = texto_linha(words)

        if texto.startswith('Proprietário:'):
          if texto != proprietario_atual:
            proprietario_atual = texto
          pending_first_line = None
          continue

        if texto.startswith('Totais - Viagens:'):
          pending_first_line = None
          continue

        if eh_linha_data(words):
          if indice > 0 and eh_possivel_primeira_linha(grupos[indice - 1][1]):
            primeira_linha = grupos[indice - 1][1]
          else:
            primeira_linha = pending_first_line

          if primeira_linha is None:
            continue

          x22 = [word['text'] for word in primeira_linha if word['x0'] < 50]
          conhec = (
              int(x22[0]) if x22 and INT_RE.fullmatch(x22[0]) else None
          )
          x22_data = [word['text'] for word in words if word['x0'] < 50][0]

          registro = {
              'Conhec.': conhec,
              'Emissão': datetime.strptime(x22_data, '%d/%m/%Y'),
              'Origem': texto_por_x(primeira_linha, 60, 170),
              'Destino': texto_por_x(words, 60, 170),
              'Motorista': texto_por_x(primeira_linha, 170, 330),
              'Mercadoria': texto_por_x(words, 170, 300),
              'Quant': valor_numerico(texto_por_x(words, 300, 338)),
              'Veículo': texto_por_x(primeira_linha, 330, 382),
              'No. NF': valor_numerico(texto_por_x(words, 338, 378)),
              'Vl. Seguro': campo_numerico(primeira_linha, 0),
              'Vl. ICMS': campo_numerico(primeira_linha, 1),
              'Peso Saída': campo_numerico(primeira_linha, 2),
              'Peso Cheg.': campo_numerico(primeira_linha, 3),
              'PPT Mot.': campo_numerico(primeira_linha, 4),
              'Fr. Emp': campo_numerico(primeira_linha, 5),
              'Base Calc.': campo_numerico(primeira_linha, 6),
              'Margem': campo_numerico(primeira_linha, 7),
              'Desp. Extra': campo_numerico(words, 0),
              'Vl. Ped.': campo_numerico(words, 1),
              'PPT Emp.': campo_numerico(words, 3),
              'Adt. Mot.': campo_numerico(words, 4),
              'Fr. Mot': campo_numerico(words, 5),
              'Fr. Fiscal': campo_numerico(words, 6),
              '% Margem': campo_numerico(words, 7),
          }
          registros.append((proprietario_atual, registro))
          pending_first_line = None
          continue

        if eh_possivel_primeira_linha(words):
          pending_first_line = words

      if not grupos or not eh_possivel_primeira_linha(grupos[-1][1]):
        pending_first_line = None

  return registros


def gerar_excel_stream(registros):
  wb = Workbook()
  ws = wb.active
  ws.title = 'Viagens'

  preenchimento_titulo = PatternFill(fill_type='solid', fgColor='D9EAF7')
  preenchimento_cabecalho = PatternFill(fill_type='solid', fgColor='B4C6E7')
  fonte_titulo = Font(bold=True, size=12)
  fonte_cabecalho = Font(bold=True, size=10)
  borda = Border(
      left=Side(style='thin', color='BFBFBF'),
      right=Side(style='thin', color='BFBFBF'),
      top=Side(style='thin', color='BFBFBF'),
      bottom=Side(style='thin', color='BFBFBF'),
  )
  alinhamento_central = Alignment(horizontal='center', vertical='center')
  alinhamento_esquerda = Alignment(horizontal='left', vertical='center')
  alinhamento_direita = Alignment(horizontal='right', vertical='center')

  linha = 1
  proprietario_anterior = None

  for proprietario, registro in registros:
    if proprietario != proprietario_anterior:
      if proprietario_anterior is not None:
        linha += 1

      ws.merge_cells(
          start_row=linha,
          start_column=1,
          end_row=linha,
          end_column=len(COLUNAS),
      )
      celula_titulo = ws.cell(row=linha, column=1, value=proprietario)
      celula_titulo.font = fonte_titulo
      celula_titulo.fill = preenchimento_titulo
      celula_titulo.alignment = alinhamento_esquerda
      linha += 1

      for coluna, nome in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=linha, column=coluna, value=nome)
        celula.font = fonte_cabecalho
        celula.fill = preenchimento_cabecalho
        celula.alignment = alinhamento_central
        celula.border = borda

      linha += 1
      proprietario_anterior = proprietario

    valores = [registro[coluna] for coluna in COLUNAS]
    for coluna, valor in enumerate(valores, start=1):
      celula = ws.cell(row=linha, column=coluna, value=valor)
      celula.border = borda

      if coluna == 2:
        celula.number_format = 'dd/mm/yyyy'
        celula.alignment = alinhamento_central
      elif coluna in (1, 9):
        celula.alignment = alinhamento_central
      elif coluna in (
          7,
          10,
          11,
          12,
          13,
          14,
          15,
          16,
          17,
          18,
          19,
          20,
          21,
          22,
          23,
          24,
      ):
        celula.alignment = alinhamento_direita
      else:
        celula.alignment = alinhamento_esquerda

    for coluna in (7, 10, 11, 12, 13, 14, 19, 20, 21, 22, 23):
      ws.cell(linha, coluna).number_format = '#,##0.00'
    for coluna in (15, 16, 17, 18):
      ws.cell(linha, coluna).number_format = '0.0000'
    ws.cell(linha, 24).number_format = '0.00'

    linha += 1

  larguras = {
      1: 11,
      2: 13,
      3: 31,
      4: 31,
      5: 37,
      6: 38,
      7: 14,
      8: 12,
      9: 12,
      10: 13,
      11: 13,
      12: 13,
      13: 13,
      14: 14,
      15: 13,
      16: 13,
      17: 13,
      18: 13,
      19: 13,
      20: 13,
      21: 13,
      22: 13,
      23: 13,
      24: 12,
  }
  for coluna, largura in larguras.items():
    ws.column_dimensions[get_column_letter(coluna)].width = largura

  #ws.freeze_panes = 'A3'
  ws.auto_filter.ref = f'A2:{get_column_letter(len(COLUNAS))}{ws.max_row}'

  # Salva o arquivo diretamente na memória RAM (em vez de gravar no HD)
  excel_output = io.BytesIO()
  wb.save(excel_output)
  excel_output.seek(0)
  return excel_output


# ROTA WEB - Página principal
@app.route('/', methods=['GET'])
def index():
  return render_template('index.html')


# ROTA WEB - Processamento do arquivo e Download
@app.route('/convert', methods=['POST'])
def convert():
  if 'pdf_file' not in request.files:
    return 'Nenhum arquivo enviado.', 400

  file = request.files['pdf_file']
  if file.filename == '':
    return 'Nenhum arquivo selecionado.', 400

  # Executa o processamento do PDF enviado
  if not file.filename.lower().endswith('.pdf'):
    return 'Envie um arquivo PDF.', 400
  try:
    registros = processar_pdf_stream(file.stream)
  except Exception:
    app.logger.exception('Falha ao ler PDF de viagens')
    return 'Não foi possível ler o PDF. Confira o arquivo enviado.', 400

  if not registros:
    return 'Nenhuma viagem foi encontrada no PDF enviado.', 400

  # Gera a planilha Excel em memória
  excel_stream = gerar_excel_stream(registros)

  # Devolve a planilha para download automático
  return send_file(
      excel_stream,
      as_attachment=True,
      download_name='Carta_Frete_Gerado.xlsx',
      mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
  )


