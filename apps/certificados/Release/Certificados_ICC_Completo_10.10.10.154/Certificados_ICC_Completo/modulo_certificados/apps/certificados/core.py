"""Inventário incremental de arquivos A1. Nenhuma chave privada é gravada no índice."""
from __future__ import annotations
import hashlib
import json
import math
import os
import sqlite3
import threading
import time
from configparser import ConfigParser
from contextlib import contextmanager
from datetime import datetime, timezone, timedelta
from pathlib import Path
from cryptography.fernet import Fernet
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.serialization import pkcs12
from cryptography.x509.oid import NameOID
from .leitura_original import extrair_cliente_e_senha, extrair_documento_certificado

UTC = timezone.utc
MAX_PFX = 10 * 1024 * 1024
LABELS = {'valido':'Válido','proximo':'Próximo do vencimento','vencido':'Vencido',
          'erro':'Falha na leitura','futuro':'Ainda não vigente'}

def now():
    return datetime.now(UTC)

def digest(value):
    return hashlib.sha256(value.encode('utf-8')).hexdigest()

def config(base):
    c = ConfigParser(interpolation=None)
    c.read([base/'config.ini', base/'config_apps.ini'], encoding='utf-8-sig')
    section = 'certificados_painel'
    def get(key, default): return c.get(section, key, fallback=default).strip()
    def path(key, default):
        p = Path(get(key, default))
        return p if p.is_absolute() else base/p
    folders = []
    for kind in ('cpf','cnpj'):
        original = c.get('monitora_certificados', 'pasta_'+kind,
                         fallback=rf'A:\CERTIFICADOS\Certificados {kind.upper()}')
        folders.append((path('pasta_'+kind, original), kind.upper()))
    return dict(base=base, folders=folders, data=base/'dados/apps/certificados',
                days=c.getint(section,'dias_proximo',fallback=30),
                interval=max(60,c.getint(section,'intervalo_segundos',fallback=300)),
                recursive=c.getboolean(section,'recursivo',fallback=False),
                state=path('estado_legado',r'C:\C\Projetos_Phyton\Certificados Digitais\Monitora Certificados\estado_certificados.json'),
                sheet=path('planilha_correcoes',r'\\srvprosoft\g\programas\Relatorio_Certificados.xlsx'),
                url=get('url_publica','').rstrip('/'))

def validade(row, today=None, days=30):
    today = today or now()
    if row.get('validacao') != 'Sucesso' or not row.get('vencimento'):
        return 'erro', None
    end = datetime.fromisoformat(row['vencimento'])
    start = datetime.fromisoformat(row['inicio'])
    remaining = (end-today).total_seconds()
    if today < start: return 'futuro', math.ceil(remaining/86400)
    if remaining < 0: return 'vencido', math.floor(remaining/86400)
    if remaining <= days*86400: return 'proximo', math.ceil(remaining/86400)
    return 'valido', math.ceil(remaining/86400)

class Inventory:
    def __init__(self, cfg):
        self.cfg = cfg
        self.directory = Path(cfg['data']); self.directory.mkdir(parents=True, exist_ok=True)
        key = self.directory/'segredo.key'
        try:
            fd = os.open(key, os.O_WRONLY|os.O_CREAT|os.O_EXCL, 0o600)
            with os.fdopen(fd,'wb') as f: f.write(Fernet.generate_key())
        except FileExistsError: pass
        self.cipher = Fernet(key.read_bytes())
        self.path = self.directory/'inventario.sqlite3'
        with self.db() as db:
            db.executescript('''
            CREATE TABLE IF NOT EXISTS certs (id TEXT PRIMARY KEY, path TEXT UNIQUE,
              signature TEXT, payload TEXT, password BLOB, manual BLOB, seen INTEGER DEFAULT 0);
            CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT);
            CREATE TABLE IF NOT EXISTS tickets (hash TEXT PRIMARY KEY, cert TEXT, uid INTEGER,
              sid_hash TEXT, route TEXT, created REAL, expires REAL, state TEXT,
              callback TEXT, callback_expires REAL, machine TEXT, sha256 TEXT, thumbprint TEXT);
            CREATE TABLE IF NOT EXISTS audit (id INTEGER PRIMARY KEY, at TEXT, uid INTEGER,
              cert TEXT, event TEXT, machine TEXT);
            ''')
        self.lock = threading.Lock()
        self.state_lock = threading.Lock()
        self.running = False; self.done = 0; self.started = 0; self.last_started = 0
        self.message = ''

    @contextmanager
    def db(self):
        con = sqlite3.connect(self.path, timeout=20)
        con.row_factory = sqlite3.Row
        try:
            with con: yield con
        finally: con.close()

    def seal(self, text): return self.cipher.encrypt(text.encode('utf-8'))
    def unseal(self, value): return self.cipher.decrypt(value).decode('utf-8') if value else None

    def audit(self, uid, cert, event, machine=''):
        with self.db() as db:
            db.execute('INSERT INTO audit(at,uid,cert,event,machine) VALUES(?,?,?,?,?)',
                       (now().isoformat(),uid,cert,event,machine[:128]))

    def start(self, force=False):
        with self.state_lock:
            if self.running or (not force and time.time()-self.last_started < self.cfg['interval']):
                return False
            self.running=True; self.done=0; self.started=time.time(); self.last_started=self.started
        threading.Thread(target=self._background,daemon=True,name='certificados-scan').start()
        return True

    def _background(self):
        try: self.scan()
        except Exception:
            self.message='Não foi possível concluir a leitura. Verifique o acesso às pastas e ao diretório de dados.'
        finally:
            with self.state_lock: self.running=False

    def _old_passwords(self):
        notes=[]; old={}; corrections={}
        try:
            if self.cfg['state'].is_file():
                raw=json.loads(self.cfg['state'].read_text(encoding='utf-8-sig'))
                if isinstance(raw,dict): old=raw
        except (OSError, ValueError): notes.append('Não foi possível ler o estado legado; usando o índice local.')
        try:
            if self.cfg['sheet'].is_file():
                from openpyxl import load_workbook
                wb=load_workbook(self.cfg['sheet'],read_only=True,data_only=True)
                try:
                    header=None
                    for i,line in enumerate(wb.active.iter_rows(values_only=True)):
                        if header is None:
                            if i>14: break
                            if 'ARQUIVO ORIGINAL' in line and 'SENHA CORRIGIDA' in line:
                                header={str(v):j for j,v in enumerate(line) if v is not None}
                            continue
                        name=line[header['ARQUIVO ORIGINAL']]; pwd=line[header['SENHA CORRIGIDA']]
                        if name and pwd is not None and str(pwd).strip():
                            if name in corrections and corrections[name] != str(pwd).strip():
                                corrections[name]=None
                            else: corrections[name]=str(pwd).strip()
                finally: wb.close()
        except Exception: notes.append('Planilha de correções indisponível; usando senhas já conhecidas.')
        return old,corrections,notes

    def _read(self, path, kind, password, client, origin):
        r=dict(tipo=kind,cliente=client,documento='',origem_senha=origin,validacao='',
               inicio=None,vencimento=None,thumbprint='',sha256='',tem_chave=False)
        if password is None:
            r['validacao']='Senha não identificada'; return r
        try:
            with path.open("rb") as source: data=source.read(MAX_PFX+1)
            if len(data)>MAX_PFX: raise ValueError('Arquivo excede o limite de 10 MB')
            key,cert,_=pkcs12.load_key_and_certificates(data,password.encode('utf-8') if password else None)
            if cert is None or key is None: raise ValueError('Sem certificado/chave privada')
            r.update(validacao='Sucesso',inicio=cert.not_valid_before_utc.isoformat(),
                     vencimento=cert.not_valid_after_utc.isoformat(),
                     thumbprint=cert.fingerprint(hashes.SHA1()).hex().upper(),
                     sha256=hashlib.sha256(data).hexdigest(),tem_chave=True)
            try: r['documento']=extrair_documento_certificado(cert,kind)
            except Exception: pass
            common=cert.subject.get_attributes_for_oid(NameOID.COMMON_NAME)
            if common: r['cliente']=common[0].value.rsplit(':',1)[0].strip() or client
        except Exception:
            r['validacao']='Senha incorreta, arquivo inacessível ou certificado inválido'
        return r

    def scan(self):
        with self.lock:
            old,corrections,notes=self._old_passwords()
            # A lista publicada nunca inclui registros que só existem no JSON antigo.
            with self.db() as db: db.execute('UPDATE certs SET seen=0')
            roots=[]; seen_paths=set()
            for root,kind in self.cfg['folders']:
                info=dict(tipo=kind,caminho=str(root),acessivel=False,quantidade=0)
                roots.append(info)
                try:
                    if not root.is_dir(): raise OSError('Pasta indisponível')
                    root_real=root.resolve(strict=True)
                    def entries(folder):
                        with os.scandir(folder) as it:
                            for e in it:
                                if e.is_symlink(): continue
                                if e.is_file(follow_symlinks=False) and Path(e.name).suffix.lower() in {'.pfx','.p12'}:
                                    yield Path(e.path)
                                elif self.cfg['recursive'] and e.is_dir(follow_symlinks=False):
                                    yield from entries(e.path)
                    for path in entries(root):
                        try:
                            resolved=path.resolve(strict=True)
                            if not resolved.is_relative_to(root_real): continue
                            canonical=os.path.normcase(str(resolved))
                            if canonical in seen_paths: continue
                            seen_paths.add(canonical)
                            stat=resolved.stat()
                            ident=digest(canonical)
                            with self.db() as db:
                                saved=db.execute('SELECT * FROM certs WHERE id=?',(ident,)).fetchone()
                            client,guessed=extrair_cliente_e_senha(path.name)
                            reg=old.get(str(path),old.get(canonical,{}))
                            pwd=self.unseal(saved['manual']) if saved and saved['manual'] else None
                            origin='Corrigida no painel'
                            if pwd is None:
                                pwd=corrections.get(path.name) or reg.get('senha_manual_salva') or None
                                origin='Manual (planilha/legado)'
                            if pwd is None and saved and saved['password']:
                                previous=json.loads(saved['payload'])
                                if previous.get('validacao')=='Sucesso' and previous.get('origem_senha') in {'Manual (planilha/legado)','Estado legado'}:
                                    pwd=self.unseal(saved['password']); origin=previous['origem_senha']
                            if pwd is None:
                                pwd=guessed if guessed != 'NÃO IDENTIFICADA' else None
                                origin='Nome do arquivo'
                                if pwd is None and reg.get('status_validacao')=='Sucesso':
                                    pwd=reg.get('senha_utilizada'); origin='Estado legado'
                            signature=f'{stat.st_mtime_ns}:{stat.st_size}:{digest(str(pwd))}'
                            previous=json.loads(saved['payload']) if saved else {}
                            if saved and saved['signature']==signature and previous.get('validacao')=='Sucesso':
                                data=previous
                            else: data=self._read(resolved,kind,pwd,client,origin)
                            data.update(id=ident,arquivo=ident[:12]+path.suffix.lower())
                            if not resolved.is_file(): continue
                            with self.db() as db:
                                db.execute('''INSERT INTO certs(id,path,signature,payload,password,seen) VALUES(?,?,?,?,?,1)
                                ON CONFLICT(id) DO UPDATE SET path=excluded.path,signature=excluded.signature,
                                payload=excluded.payload,password=excluded.password,seen=1''',
                                (ident,str(resolved),signature,json.dumps(data,ensure_ascii=False),self.seal(pwd) if pwd is not None else None))
                            self.done+=1; info['quantidade']+=1
                        except (OSError, ValueError):
                            notes.append(f'Um arquivo {kind} mudou ou ficou indisponível durante a leitura.')
                    info['acessivel']=True
                except OSError:
                    notes.append(f'Pasta {kind} indisponível ou leitura incompleta. Confira o caminho e a conta que executa o portal.')
            meta=dict(atualizado_em=now().isoformat(),pastas=roots,avisos=list(dict.fromkeys(notes)))
            with self.db() as db:
                db.execute("INSERT OR REPLACE INTO meta VALUES('scan',?)",(json.dumps(meta,ensure_ascii=False),))
            self.message=''

    def snapshot(self):
        with self.db() as db:
            rows=[json.loads(r[0]) for r in db.execute('SELECT payload FROM certs WHERE seen=1')]
            meta=db.execute("SELECT value FROM meta WHERE key='scan'").fetchone()
        counts={k:0 for k in LABELS}
        for row in rows:
            status,days=validade(row,days=self.cfg['days'])
            row.update(status=status,status_label=LABELS[status],dias=days,
                       instalavel=status in {'valido','proximo'} and row['tem_chave'])
            counts[status]+=1
        metadata=json.loads(meta[0]) if meta else dict(atualizado_em=None,pastas=[],avisos=[])
        metadata.update(em_andamento=self.running,processados=self.done,mensagem=self.message)
        return rows,counts,metadata

    def get_file(self, ident):
        with self.db() as db: row=db.execute('SELECT * FROM certs WHERE id=? AND seen=1',(ident,)).fetchone()
        if not row: raise ValueError('Certificado não encontrado. Atualize a lista.')
        p=Path(row['path']).resolve(strict=True)
        if not any(p.is_relative_to(root.resolve(strict=True)) for root,_ in self.cfg['folders'] if root.is_dir()):
            raise ValueError('Arquivo fora das pastas autorizadas.')
        if not p.is_file() or p.stat().st_size>MAX_PFX: raise ValueError('Arquivo indisponível.')
        return row,p,json.loads(row['payload'])

    def correct(self, ident, password):
        with self.lock:
            row,p,old=self.get_file(ident)
            data=self._read(p,old['tipo'],password,old['cliente'],'Corrigida no painel')
            if data['validacao']!='Sucesso': raise ValueError('A senha não abriu este certificado. Nenhuma alteração foi salva.')
            data.update(id=ident,arquivo=old['arquivo'])
            st=p.stat(); signature=f'{st.st_mtime_ns}:{st.st_size}:{digest(str(password))}'
            with self.db() as db:
                db.execute('UPDATE certs SET manual=?,password=?,payload=?,signature=? WHERE id=?',
                           (self.seal(password),self.seal(password),json.dumps(data,ensure_ascii=False),signature,ident))
