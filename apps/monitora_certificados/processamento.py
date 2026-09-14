import os
import re
import json
import pandas as pd
from datetime import datetime
from cryptography.hazmat.primitives.serialization import pkcs12
from cryptography.hazmat.backends import default_backend
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLUNAS_RELATORIO = [
    "TIPO",
    "NOME DO CLIENTE",
    "ARQUIVO ORIGINAL",
    "SENHA UTILIZADA",
    "ORIGEM DA SENHA",
    "SENHA CORRIGIDA",
    "STATUS DA VALIDAÇÃO",
    "VENCIMENTO REAL",
    "DIAS RESTANTES",
    "STATUS DE VALIDADE"
]


def formatar_planilha(caminho_planilha):
    workbook = load_workbook(caminho_planilha)
    worksheet = workbook.active
    worksheet.title = "Certificados"

    # Reserva as primeiras linhas para o dashboard. O cabeçalho original,
    # criado pelo pandas, passa da linha 1 para a linha 7.
    worksheet.insert_rows(1, amount=6)
    linha_cabecalho = 7
    primeira_linha_dados = linha_cabecalho + 1

    azul_claro = "9DC3E6"
    azul_escuro = "1F4E78"
    cinza_claro = "E7E6E6"
    branco = "FFFFFF"
    verde_claro = "C6EFCE"
    verde_escuro = "006100"
    amarelo_claro = "FFEB9C"
    amarelo_escuro = "9C6500"
    vermelho_claro = "FFC7CE"
    vermelho_escuro = "9C0006"
    cinza_status = "D9E1F2"

    borda_fina = Side(style="thin", color="B7B7B7")
    borda = Border(
        left=borda_fina,
        right=borda_fina,
        top=borda_fina,
        bottom=borda_fina
    )

    # ---------------------------------------------------------------------
    # DASHBOARD
    # ---------------------------------------------------------------------
    worksheet.merge_cells("A1:J1")
    titulo_dashboard = worksheet["A1"]
    titulo_dashboard.value = "RESUMO DOS CERTIFICADOS DIGITAIS"
    titulo_dashboard.fill = PatternFill("solid", fgColor=azul_escuro)
    titulo_dashboard.font = Font(bold=True, color=branco, size=14)
    titulo_dashboard.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 28

    coluna_status_dashboard = 10  # Coluna J: STATUS DE VALIDADE
    status_encontrados = [
        worksheet.cell(linha, coluna_status_dashboard).value
        for linha in range(primeira_linha_dados, worksheet.max_row + 1)
    ]
    total_vencidos = status_encontrados.count("VENCIDO")
    total_a_vencer = status_encontrados.count("VENCE EM BREVE (<30 dias)")
    total_validos = status_encontrados.count("VÁLIDO")

    cartoes = [
        {
            "intervalo": "A2:C5",
            "titulo": "CERTIFICADOS VENCIDOS",
            "quantidade": total_vencidos,
            "descricao": "Indisponíveis para uso",
            "fundo": vermelho_claro,
            "fonte": vermelho_escuro,
        },
        {
            "intervalo": "D2:G5",
            "titulo": "A VENCER OU PRÓXIMOS DO VENCIMENTO",
            "quantidade": total_a_vencer,
            "descricao": "Vencem em até 30 dias",
            "fundo": amarelo_claro,
            "fonte": amarelo_escuro,
        },
        {
            "intervalo": "H2:J5",
            "titulo": "CERTIFICADOS VÁLIDOS",
            "quantidade": total_validos,
            "descricao": "Disponíveis para uso",
            "fundo": verde_claro,
            "fonte": verde_escuro,
        },
    ]

    for cartao in cartoes:
        inicio, fim = cartao["intervalo"].split(":")
        coluna_inicial = worksheet[inicio].column
        coluna_final = worksheet[fim].column
        linha_inicial = worksheet[inicio].row
        linha_final = worksheet[fim].row

        # Aplica fundo e bordas antes de mesclar, mantendo o contorno do cartão.
        for linha in worksheet.iter_rows(
            min_row=linha_inicial,
            max_row=linha_final,
            min_col=coluna_inicial,
            max_col=coluna_final,
        ):
            for celula in linha:
                celula.fill = PatternFill("solid", fgColor=cartao["fundo"])
                celula.border = borda

        faixa_titulo = (
            f"{get_column_letter(coluna_inicial)}{linha_inicial}:"
            f"{get_column_letter(coluna_final)}{linha_inicial}"
        )
        faixa_numero = (
            f"{get_column_letter(coluna_inicial)}{linha_inicial + 1}:"
            f"{get_column_letter(coluna_final)}{linha_inicial + 2}"
        )
        faixa_descricao = (
            f"{get_column_letter(coluna_inicial)}{linha_final}:"
            f"{get_column_letter(coluna_final)}{linha_final}"
        )
        worksheet.merge_cells(faixa_titulo)
        worksheet.merge_cells(faixa_numero)
        worksheet.merge_cells(faixa_descricao)

        celula_titulo = worksheet.cell(linha_inicial, coluna_inicial)
        celula_titulo.value = cartao["titulo"]
        celula_titulo.font = Font(bold=True, color=cartao["fonte"], size=10)
        celula_titulo.alignment = Alignment(horizontal="center", vertical="center")

        celula_numero = worksheet.cell(linha_inicial + 1, coluna_inicial)
        celula_numero.value = cartao["quantidade"]
        celula_numero.font = Font(bold=True, color=cartao["fonte"], size=22)
        celula_numero.alignment = Alignment(horizontal="center", vertical="center")

        celula_descricao = worksheet.cell(linha_final, coluna_inicial)
        celula_descricao.value = cartao["descricao"]
        celula_descricao.font = Font(color=cartao["fonte"], size=9)
        celula_descricao.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.row_dimensions[2].height = 24
    worksheet.row_dimensions[3].height = 24
    worksheet.row_dimensions[4].height = 24
    worksheet.row_dimensions[5].height = 22
    worksheet.row_dimensions[6].height = 8

    # ---------------------------------------------------------------------
    # TABELA DE CERTIFICADOS
    # ---------------------------------------------------------------------
    worksheet.row_dimensions[linha_cabecalho].height = 30
    for celula in worksheet[linha_cabecalho]:
        celula.fill = PatternFill("solid", fgColor=azul_claro)
        celula.font = Font(bold=True, color=azul_escuro)
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        celula.border = borda

    for numero_linha, linha in enumerate(
        worksheet.iter_rows(min_row=primeira_linha_dados, max_row=worksheet.max_row),
        start=primeira_linha_dados
    ):
        cor_linha = cinza_claro if numero_linha % 2 == 0 else branco
        for celula in linha:
            celula.fill = PatternFill("solid", fgColor=cor_linha)
            celula.border = borda
            celula.alignment = Alignment(vertical="center")

    cabecalhos = {
        celula.value: celula.column
        for celula in worksheet[linha_cabecalho]
        if celula.value is not None
    }

    for nome_coluna in ("SENHA UTILIZADA", "SENHA CORRIGIDA"):
        coluna = cabecalhos.get(nome_coluna)
        if coluna:
            for celula in worksheet.iter_cols(
                min_col=coluna,
                max_col=coluna,
                min_row=primeira_linha_dados,
                max_row=worksheet.max_row
            ):
                celula[0].number_format = "@"

    coluna_vencimento = cabecalhos.get("VENCIMENTO REAL")
    if coluna_vencimento:
        for linha in range(primeira_linha_dados, worksheet.max_row + 1):
            celula = worksheet.cell(linha, coluna_vencimento)
            if isinstance(celula.value, str):
                try:
                    celula.value = datetime.strptime(
                        celula.value,
                        "%Y-%m-%d %H:%M:%S"
                    )
                    celula.number_format = "dd/mm/yyyy hh:mm:ss"
                except ValueError:
                    pass
            celula.alignment = Alignment(horizontal="center", vertical="center")

    coluna_dias = cabecalhos.get("DIAS RESTANTES")
    if coluna_dias:
        for linha in range(primeira_linha_dados, worksheet.max_row + 1):
            celula = worksheet.cell(linha, coluna_dias)
            celula.number_format = "0"
            celula.alignment = Alignment(horizontal="center", vertical="center")

    coluna_status = cabecalhos.get("STATUS DE VALIDADE")
    if coluna_status:
        cores_status = {
            "VÁLIDO": (verde_claro, verde_escuro),
            "VENCE EM BREVE (<30 dias)": (amarelo_claro, amarelo_escuro),
            "VENCIDO": (vermelho_claro, vermelho_escuro),
            "DESCONHECIDO": (cinza_status, azul_escuro)
        }
        for linha in range(primeira_linha_dados, worksheet.max_row + 1):
            celula = worksheet.cell(linha, coluna_status)
            preenchimento, cor_fonte = cores_status.get(
                celula.value,
                (cinza_status, azul_escuro)
            )
            celula.fill = PatternFill("solid", fgColor=preenchimento)
            celula.font = Font(bold=True, color=cor_fonte)
            celula.alignment = Alignment(horizontal="center", vertical="center")

    larguras_minimas = {
        "TIPO": 10,
        "NOME DO CLIENTE": 28,
        "ARQUIVO ORIGINAL": 42,
        "SENHA UTILIZADA": 20,
        "ORIGEM DA SENHA": 20,
        "SENHA CORRIGIDA": 20,
        "STATUS DA VALIDAÇÃO": 42,
        "VENCIMENTO REAL": 23,
        "DIAS RESTANTES": 16,
        "STATUS DE VALIDADE": 29
    }

    for indice_coluna in range(1, worksheet.max_column + 1):
        letra_coluna = get_column_letter(indice_coluna)
        nome_coluna = worksheet.cell(linha_cabecalho, indice_coluna).value
        maior_conteudo = max(
            len(str(worksheet.cell(linha, indice_coluna).value or ""))
            for linha in range(linha_cabecalho, worksheet.max_row + 1)
        )
        largura_calculada = maior_conteudo + 2
        largura_minima = larguras_minimas.get(nome_coluna, 12)
        worksheet.column_dimensions[letra_coluna].width = min(
            max(largura_calculada, largura_minima),
            55
        )

    worksheet.freeze_panes = f"A{primeira_linha_dados}"
    worksheet.sheet_view.showGridLines = False
    worksheet.auto_filter.ref = (
        f"A{linha_cabecalho}:"
        f"{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    )

    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = f"1:{linha_cabecalho}"

    workbook.save(caminho_planilha)


def extrair_cliente_e_senha(nome_arquivo):
    nome = re.sub(r'(?i)(\.pfx|\.p12)+$', '', nome_arquivo)
    nome = nome.strip()
    
    nome = re.sub(r'\s*\(\d+\)$', '', nome)
    nome = re.sub(r'\(ARROBA\)$', '', nome, flags=re.IGNORECASE)
    nome = re.sub(r'[\s\-]*CPF(?:\s+novo)?$', '', nome, flags=re.IGNORECASE)
    nome = re.sub(r'_VENC\..*$', '', nome, flags=re.IGNORECASE)
    nome = re.sub(r'\s+-$', '', nome)
    
    senha = ""
    cliente = ""
    
    m = re.search(r"^(.*?)\s+\d{11,14}\s+\d{2}-\d{2}-\d{4}\s+(.*)$", nome)
    if m: 
        cliente, senha = m.group(1), m.group(2)
    else:
        m = re.search(r"^(.*?)(?:[_\s-]*senha[_\s-]+)(.*)$", nome, re.IGNORECASE)
        if m: 
            cliente, senha = m.group(1), m.group(2)
        else:
            m = re.search(r"^(.*?)[_\s-]+(\d{3,})$", nome)
            if m:
                cliente, senha = m.group(1), m.group(2)
            else:
                m = re.search(r"^(.*?)\s*-\s*(.*)$", nome)
                if m: 
                    cliente, senha = m.group(1), m.group(2)
                else:
                    m = re.search(r"^(.*?)_([^_]+)$", nome)
                    if m and (re.search(r"\d", m.group(2))):
                        cliente, senha = m.group(1), m.group(2)
                    else:
                        m = re.search(r"^(.*?)\s+([^_\s]+)$", nome)
                        if m and (re.search(r"\d", m.group(2))):
                            cliente, senha = m.group(1), m.group(2)
                        else:
                            cliente, senha = nome, "NÃO IDENTIFICADA"
                            
    if senha != "NÃO IDENTIFICADA":
        cliente = cliente.strip(' _-')
        senha_limpa = senha.lstrip('_- ')
        
        if re.match(r'^\d+$', senha_limpa):
            senha = senha_limpa
        else:
            senha = senha.lstrip('_')
            
    return cliente, senha


def obter_vencimento_real(caminho_arquivo, senha):
    try:
        with open(caminho_arquivo, "rb") as f:
            pfx_data = f.read()
            
        _, certificate, _ = pkcs12.load_key_and_certificates(
            pfx_data,
            senha.encode('utf-8'),
            backend=default_backend()
        )
        return certificate.not_valid_after
    except Exception:
        return None


def carregar_senhas_manuais(caminho_planilha):
    senhas_manuais = {}
    if os.path.exists(caminho_planilha):
        try:
            # Localiza o cabeçalho para aceitar tanto as planilhas antigas,
            # iniciadas na linha 1, quanto as novas, com dashboard no topo.
            previa = pd.read_excel(caminho_planilha, header=None, nrows=15, dtype=str)
            linha_cabecalho = 0
            for indice, linha in previa.iterrows():
                valores = {str(valor).strip() for valor in linha.dropna().tolist()}
                if "ARQUIVO ORIGINAL" in valores and "SENHA CORRIGIDA" in valores:
                    linha_cabecalho = indice
                    break

            df_antigo = pd.read_excel(
                caminho_planilha,
                header=linha_cabecalho,
                dtype=str,
            )
            if "SENHA CORRIGIDA" in df_antigo.columns:
                for _, linha in df_antigo.iterrows():
                    senha_preenchida = linha.get("SENHA CORRIGIDA")
                    if pd.notna(senha_preenchida):
                        senha_str = str(senha_preenchida).strip()
                        if senha_str != "" and senha_str.lower() != 'nan':
                            senhas_manuais[linha["ARQUIVO ORIGINAL"]] = senha_str
        except Exception as e:
            print(f"Aviso: Não foi possível ler a planilha existente. Detalhe: {e}")
    return senhas_manuais


def validar_certificados(pastas_alvo, pasta_projeto, pasta_planilha):
    
    # 1. Diretório do Projeto (Memória JSON)
    if not os.path.exists(pasta_projeto):
        os.makedirs(pasta_projeto)
    caminho_estado = os.path.join(pasta_projeto, "estado_certificados.json")

    # 2. Diretório da Rede (Relatório Excel)
    if not os.path.exists(pasta_planilha):
        os.makedirs(pasta_planilha)
    caminho_planilha = os.path.join(pasta_planilha, "Relatorio_Certificados.xlsx")
    
    estado_atual = {}
    if os.path.exists(caminho_estado):
        with open(caminho_estado, "r", encoding="utf-8") as f:
            estado_atual = json.load(f)

    senhas_manuais = carregar_senhas_manuais(caminho_planilha)
    hoje = datetime.now()
    dados_planilha = []
    arquivos_atuais = set()
    pastas_verificadas = set()

    for pasta, tipo in pastas_alvo.items():
        if not os.path.exists(pasta):
            print(f"Aviso: A pasta '{pasta}' não foi encontrada.")
            continue

        pasta_normalizada = os.path.normcase(os.path.normpath(pasta))
        pastas_verificadas.add(pasta_normalizada)

        # A listagem é capturada novamente em cada execução. Assim, somente os
        # certificados que existem na pasta neste momento entram no relatório.
        for entrada in os.scandir(pasta):
            if not entrada.is_file() or not entrada.name.lower().endswith(('.pfx', '.p12')):
                continue

            nome_arquivo = entrada.name
            caminho_arquivo = entrada.path

            # Evita incluir um arquivo removido durante a própria execução.
            if not os.path.isfile(caminho_arquivo):
                continue

            timestamp_modificacao = os.path.getmtime(caminho_arquivo)
            id_arquivo = os.path.join(pasta, nome_arquivo)
            arquivos_atuais.add(os.path.normcase(os.path.normpath(id_arquivo)))
            
            registro = estado_atual.get(id_arquivo, {})
            
            senha_manual = senhas_manuais.get(nome_arquivo)
            if not senha_manual:
                senha_manual = registro.get("senha_manual_salva", "")

            if senha_manual:
                nome_cliente, _ = extrair_cliente_e_senha(nome_arquivo)
                senha_cert = senha_manual
                origem_senha = "Manual"
            else:
                nome_cliente, senha_cert = extrair_cliente_e_senha(nome_arquivo)
                origem_senha = "Automática"
            
            precisa_validar = False
            ultima_senha_testada = registro.get("senha_utilizada", "")
            
            if (not registro or 
                registro.get("modificado_em") != timestamp_modificacao or 
                registro.get("status_validacao") != "Sucesso" or 
                senha_cert != ultima_senha_testada):
                precisa_validar = True

            vencimento_str = registro.get("vencimento")
            status_validacao = registro.get("status_validacao", "Pendente")
            
            if precisa_validar:
                if senha_cert == "NÃO IDENTIFICADA":
                    status_validacao = "Erro: Senha não identificada"
                else:
                    data_vencimento = obter_vencimento_real(caminho_arquivo, senha_cert)
                    if data_vencimento:
                        vencimento_str = data_vencimento.strftime("%Y-%m-%d %H:%M:%S")
                        status_validacao = "Sucesso"
                    else:
                        status_validacao = "Erro: Senha incorreta ou arquivo corrompido"
                        vencimento_str = None
                
                estado_atual[id_arquivo] = {
                    "modificado_em": timestamp_modificacao,
                    "vencimento": vencimento_str,
                    "status_validacao": status_validacao,
                    "senha_utilizada": senha_cert,
                    "senha_manual_salva": senha_manual 
                }
            
            status_validade = "DESCONHECIDO"
            dias_restantes = ""
            
            if vencimento_str:
                data_venc = datetime.strptime(vencimento_str, "%Y-%m-%d %H:%M:%S")
                dias_restantes = (data_venc - hoje).days
                
                if dias_restantes < 0:
                    status_validade = "VENCIDO"
                elif dias_restantes <= 30:
                    status_validade = "VENCE EM BREVE (<30 dias)"
                else:
                    status_validade = "VÁLIDO"

            dados_planilha.append({
                "TIPO": tipo,
                "NOME DO CLIENTE": nome_cliente,
                "ARQUIVO ORIGINAL": nome_arquivo,
                "SENHA UTILIZADA": senha_cert,
                "ORIGEM DA SENHA": origem_senha,
                "SENHA CORRIGIDA": senha_manual,
                "STATUS DA VALIDAÇÃO": status_validacao,
                "VENCIMENTO REAL": vencimento_str if vencimento_str else "Falha na Leitura",
                "DIAS RESTANTES": dias_restantes,
                "STATUS DE VALIDADE": status_validade
            })

    # Remove da memória JSON os certificados que já não existem nas pastas
    # verificadas. Registros de uma unidade indisponível não são apagados.
    for id_salvo in list(estado_atual):
        id_normalizado = os.path.normcase(os.path.normpath(id_salvo))
        pasta_id = os.path.normcase(os.path.normpath(os.path.dirname(id_salvo)))
        if pasta_id in pastas_verificadas and id_normalizado not in arquivos_atuais:
            del estado_atual[id_salvo]

    with open(caminho_estado, "w", encoding="utf-8") as f:
        json.dump(estado_atual, f, indent=4, ensure_ascii=False)

    # A planilha é sempre recriada, mesmo quando não houver certificados. Isso
    # impede que linhas de uma execução anterior permaneçam no relatório.
    df = pd.DataFrame(dados_planilha, columns=COLUNAS_RELATORIO)
    df.to_excel(caminho_planilha, index=False)
    formatar_planilha(caminho_planilha)

    print(
        f"Processo concluído. {len(dados_planilha)} certificado(s) "
        f"encontrado(s). Planilha atualizada em:\n{caminho_planilha}"
    )

