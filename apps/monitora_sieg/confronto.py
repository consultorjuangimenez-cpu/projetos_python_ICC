"""Leitura somente consulta e regras do confronto SIEG x ICC."""
from __future__ import annotations

import re
import logging
import unicodedata
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


def texto(value):
    return "" if value is None else str(value).strip()


def normalizar(value):
    value = unicodedata.normalize("NFKD", texto(value))
    return " ".join("".join(c for c in value if not unicodedata.combining(c)).upper().split())


def documento(value, tipo=""):
    """Não converte CPF/CNPJ para float. Recupera zeros apenas com tipo conhecido."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            return ""
        value = str(int(value))
    value = texto(value)
    if not re.fullmatch(r"[\d./\-\s]+", value):
        return ""
    value = re.sub(r"\D", "", value)
    if not value or len(value) > 14:
        return ""
    tipo = normalizar(tipo)
    if tipo in ("CPF", "CNPJ"):
        size = 11 if tipo == "CPF" else 14
        return value.zfill(size) if len(value) <= size else ""
    return value


def data_iso(value, epoch=None):
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return datetime.combine(value, time()).isoformat(timespec="seconds")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = from_excel(value, epoch=epoch) if epoch else from_excel(value)
            return parsed.isoformat(timespec="seconds") if isinstance(parsed, datetime) else ""
        except (ValueError, OverflowError):
            return ""
    raw = texto(value)
    for pattern in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, pattern).isoformat(timespec="seconds")
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(raw).replace(tzinfo=None).isoformat(timespec="seconds")
    except ValueError:
        return ""


VERSAO_BUSCA = "MODIFICACAO-2026-09-10-R2"
_LOG = logging.getLogger(__name__)
_LOG.warning("Monitor SIEG | %s | modulo=%s", VERSAO_BUSCA, Path(__file__).resolve())


def ultima_planilha(pasta, prefixo):
    """Seleciona pela modificação; a data no nome é apenas informativa."""
    pasta = Path(pasta)
    diagnostico = f"[{VERSAO_BUSCA}] Módulo: {Path(__file__).resolve()}. Pasta consultada: {pasta}."
    if not pasta.is_dir():
        raise ValueError(f"{diagnostico} Pasta indisponível para o processo do portal. Confira o caminho e o acesso às unidades de rede.")
    encontrados = []
    try:
        itens = list(pasta.iterdir())
    except OSError as exc:
        raise ValueError(f"{diagnostico} Falha ao listar a pasta: {exc}") from exc
    for path in itens:
        if (path.name.startswith("~$") or not path.is_file()
                or path.suffix.lower() != ".xlsx"
                or not path.name.lower().startswith(prefixo.lower() + "-")):
            continue
        encontrados.append((path.stat().st_mtime_ns, path.name, path))
    if not encontrados:
        nomes = "; ".join(repr(p.name) for p in sorted(itens, key=lambda p: p.name)[:30]) or "(pasta vazia)"
        raise ValueError(f"{diagnostico} Nenhuma planilha {prefixo}-*.xlsx encontrada. Itens visíveis ({len(itens)}; até 30 exibidos): {nomes}")
    _, _, path = max(encontrados)
    _LOG.warning("Monitor SIEG | %s | selecionado=%s | modificado=%s | candidatos=%s",
                 VERSAO_BUSCA, path, datetime.fromtimestamp(path.stat().st_mtime).isoformat(), len(encontrados))
    export_date = None
    match = re.match(re.escape(prefixo) + r"-(\d{2})-(\d{2})-(\d{4}|\d{2})(?!\d)", path.name, re.I)
    if match:
        day, month, year = map(int, match.groups())
        year = 2000 + year if year < 100 else year
        try:
            export_date = date(year, month, day)
        except ValueError:
            pass
    return path, export_date


def ler_linhas(path, layout):
    """Localiza cabeçalho inclusive abaixo do dashboard ICC; exige colunas esperadas."""
    width = {"cadastros": 24, "icc": 11, "certidoes": 7}[layout]
    doc_col = 2 if layout == "icc" else 1
    status_col = {"cadastros": 14, "icc": 10, "certidoes": 5}[layout]
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        candidatas = []
        for ws in wb.worksheets:
            for number, row in enumerate(ws.iter_rows(min_row=1, max_row=100, max_col=width, values_only=True), 1):
                doc = normalizar(row[doc_col])
                status = normalizar(row[status_col])
                doc_ok = "CNPJ" in doc or "CPF" in doc
                status_ok = ("STATUS" in status or "VALIDADE" in status) if layout != "certidoes" else "SITUA" in status
                if doc_ok and status_ok:
                    candidatas.append((ws, number))
                    break
        if len(candidatas) != 1:
            reason = "Não encontrei o cabeçalho nas colunas informadas" if not candidatas else "Há mais de uma aba com o mesmo layout"
            raise ValueError(f"{path.name}: {reason}. Mantenha uma aba de dados correspondente ao relatório; o cabeçalho deve estar nas primeiras 100 linhas.")
        ws, header = candidatas[0]
        # Alguns exportadores declaram dimensões menores do que o conteúdo real.
        ws.reset_dimensions()
        linhas = []
        for number, row in enumerate(ws.iter_rows(min_row=header + 1, max_col=width, values_only=True), header + 1):
            if any(v is not None and texto(v) != "" for v in row):
                linhas.append((number, row))
        return linhas, ws.title, wb.epoch
    finally:
        wb.close()


def meta(path, aba, export_date=None):
    result = {"arquivo": path.name, "caminho": str(path), "aba": aba,
              "modificado": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")}
    if export_date:
        result["data_exportacao"] = export_date.isoformat()
    return result


def status_icc(value):
    value = normalizar(value)
    if value == "VALIDO":
        return "valido"
    if value == "VENCE EM BREVE (<30 DIAS)":
        return "breve"
    if value == "VENCIDO":
        return "vencido"
    if value == "DESCONHECIDO":
        return "desconhecido"
    return "revisar"


def carregar_icc(path):
    rows, aba, epoch = ler_linhas(path, "icc")
    index, avisos = defaultdict(list), []
    for number, r in rows:
        if normalizar(r[0]) not in ("CPF", "CNPJ"):
            if r[2] is not None:
                avisos.append(f"ICC, linha {number}: tipo diferente de CPF/CNPJ. Linha desconsiderada.")
            continue
        doc = documento(r[2], r[0])
        if not doc:
            avisos.append(f"ICC, linha {number}: CNPJ/CPF ausente ou ilegível. Linha desconsiderada.")
            continue
        index[doc].append({"documento": doc, "tipo": texto(r[0]), "nome": texto(r[1]),
                           "vencimento": data_iso(r[8], epoch), "status": texto(r[10]) or "Não informado",
                           "classe": status_icc(r[10]), "linha": number})
    if avisos and not index:
        raise ValueError(f"{path.name}: nenhum certificado com tipo e documento legíveis. Confira as colunas A (CPF/CNPJ) e C (documento).")
    return index, meta(path, aba), avisos


def resolver_documento(doc, conhecidos):
    # Um Excel numérico pode ter perdido zeros. Só recupera por correspondência única.
    candidatos = {d for d in conhecidos if d.lstrip("0") == doc.lstrip("0")} if doc else set()
    if len(candidatos) > 1:
        return doc, "Correspondência ambígua de CNPJ/CPF. Confira os zeros à esquerda nas planilhas."
    if candidatos:
        return next(iter(candidatos)), ""
    return doc, "" if len(doc) in (11, 14) else "CNPJ/CPF ausente, incompleto ou ilegível."


def carregar_cadastros(path, export_date, conhecidos):
    rows, aba, epoch = ler_linhas(path, "cadastros")
    ativos, avisos, vistos = [], [], set()
    inativos = 0
    for number, r in rows:
        if normalizar(r[0]) != "SIM":
            inativos += 1
            continue
        doc, problema = resolver_documento(documento(r[1]), conhecidos)
        raw = {"documento": doc, "nome": texto(r[3]) or texto(r[2]) or texto(r[4]) or "Nome não informado",
               "fantasia": texto(r[2]), "razao_social": texto(r[3]), "apelido": texto(r[4]),
               "estado": texto(r[5]), "vencimento_sieg": data_iso(r[13], epoch),
               "status_sieg": texto(r[14]) or "Não informado", "hub": texto(r[22]), "iris": texto(r[23]),
               "linha": number, "problema_documento": problema}
        # Duplicatas exatas não inflam contadores; registros conflitantes permanecem visíveis.
        fingerprint = tuple((k, v) for k, v in raw.items() if k != "linha")
        if fingerprint in vistos:
            avisos.append(f"Cadastros, linha {number}: cadastro ativo duplicado idêntico desconsiderado.")
            continue
        vistos.add(fingerprint)
        ativos.append(raw)
    docs = defaultdict(list)
    for row in ativos:
        if row["documento"]:
            docs[row["documento"]].append(row)
    for doc, group in docs.items():
        if len(group) > 1:
            for row in group:
                row["cadastro_duplicado"] = True
            avisos.append(f"CNPJ/CPF {doc}: mais de um cadastro ativo com informações diferentes. Confira as linhas na tela.")
    info = meta(path, aba, export_date)
    info["linhas_inativas_ignoradas"] = inativos
    return ativos, info, avisos


def comparar(cliente, certificados, hoje=None):
    hoje = hoje or date.today()
    c = dict(cliente)
    ranking = {"valido": 2, "breve": 2, "vencido": 1, "desconhecido": 0, "revisar": 0}
    ordenados = sorted(certificados, key=lambda x: (ranking[x["classe"]], x["vencimento"], -x["linha"]), reverse=True)
    selecionado = ordenados[0] if ordenados else None
    c.update({"certificados": ordenados, "certificado_icc": selecionado, "alertas": [], "proximo_vencimento": False})
    status_sieg = normalizar(c["status_sieg"])
    c["estado_pendente"] = not bool(c["estado"])
    if c["estado_pendente"]:
        c["alertas"].append("Preencher estado no cadastro SIEG")
    if c.get("cadastro_duplicado"):
        c["alertas"].append("Cadastros ativos conflitantes para o mesmo documento")
    if c["problema_documento"]:
        c.update(acao="REVISAR CNPJ/CPF", grupo="revisar", condicao=c["problema_documento"])
    elif selecionado is None:
        c.update(acao="SEM CERTIFICADO NA PASTA", grupo="sem_pasta", condicao="Documento não localizado no relatório ICC")
    elif selecionado["classe"] in ("vencido", "desconhecido"):
        c.update(acao="ATUALIZAR CERTIFICADO NA PASTA", grupo="pasta", condicao="CERTIFICADO VENCIDO OU INVÁLIDO NA PASTA")
    elif selecionado["classe"] == "revisar":
        c.update(acao="REVISAR STATUS NA PASTA", grupo="revisar", condicao="Status ICC não reconhecido")
    elif status_sieg in ("SEM CERTIFICADO", "VENCIDO"):
        c.update(acao="ATUALIZAR CERTIFICADO NO SIEG", grupo="sieg", condicao="Certificado válido disponível no relatório ICC")
    elif status_sieg == "PROXIMO DO VENCIMENTO":
        c.update(acao="ACOMPANHAR VENCIMENTO NO SIEG", grupo="acompanhar", condicao="Certificado SIEG próximo do vencimento")
    elif status_sieg == "VALIDO":
        c.update(acao="SEM PENDÊNCIA DE CERTIFICADO", grupo="ok", condicao="Certificados válidos no SIEG e no relatório ICC")
    else:
        c.update(acao="REVISAR STATUS NO SIEG", grupo="revisar", condicao="Status SIEG não reconhecido")
    if selecionado:
        c["proximo_vencimento"] = selecionado["classe"] == "breve"
        if c["proximo_vencimento"]:
            c["alertas"].append("Certificado na pasta próximo do vencimento (<30 dias)")
        venc = selecionado["vencimento"]
        if not venc:
            c["alertas"].append("Vencimento real ICC ausente ou ilegível")
        else:
            dias = (date.fromisoformat(venc[:10]) - hoje).days
            c["dias_icc"] = dias
            if (dias < 0 and selecionado["classe"] in ("valido", "breve")) or (dias > 0 and selecionado["classe"] == "vencido"):
                c["alertas"].append("Status ICC contradiz a data de vencimento. Atualize o relatório antes de agir")
            elif 0 <= dias < 30 and selecionado["classe"] == "valido":
                c["alertas"].append("Data ICC a menos de 30 dias, embora a coluna K informe VÁLIDO")
        if len(ordenados) > 1:
            c["alertas"].append(f"{len(ordenados)} certificados ICC para este documento; veja os detalhes")
    if c["vencimento_sieg"] and c["vencimento_sieg"][:10] < hoje.isoformat() and status_sieg in ("VALIDO", "PROXIMO DO VENCIMENTO"):
        c["alertas"].append("Status SIEG contradiz o vencimento. Confira a exportação")
    return c


def carregar_certidoes(path, export_date, ativos):
    rows, aba, epoch = ler_linhas(path, "certidoes")
    docs = {r["documento"] for r in ativos if not r["problema_documento"]}
    result, avisos = [], []
    ignoradas = 0
    for number, r in rows:
        doc, problema = resolver_documento(documento(r[1]), docs)
        if problema or doc not in docs:
            ignoradas += 1
            if problema:
                avisos.append(f"Certidões, linha {number}: {problema}")
            continue
        result.append({"documento": doc, "nome": texto(r[0]), "tipo": texto(r[4]) or "Não informado",
                       "situacao": texto(r[5]) or "Não informado", "vencimento": data_iso(r[6], epoch),
                       "vencimento_original": texto(r[6]), "linha": number})
        if r[6] is not None and texto(r[6]) and not result[-1]["vencimento"]:
            avisos.append(f"Certidões, linha {number}: data de vencimento ilegível ({texto(r[6])}).")
    info = meta(path, aba, export_date)
    info["linhas_sem_vinculo_ativo"] = ignoradas
    return result, info, avisos


def montar_painel(pasta_sieg, relatorio_icc):
    out = {"certificados": [], "certidoes": [], "fontes": {}, "erros": {}, "avisos": [],
           "atualizado_em": datetime.now().isoformat(timespec="seconds"), "hoje": date.today().isoformat()}
    index, ativos = None, None
    try:
        index, info, avisos = carregar_icc(Path(relatorio_icc))
        out["fontes"]["icc"] = info
        out["avisos"].extend(avisos)
    except Exception as exc:
        out["erros"]["icc"] = f"Não foi possível ler o relatório ICC: {exc}"
    try:
        path, export_date = ultima_planilha(pasta_sieg, "Cadastros")
        ativos, info, avisos = carregar_cadastros(path, export_date, index or {})
        out["fontes"]["cadastros"] = info
        out["avisos"].extend(avisos)
    except Exception as exc:
        out["erros"]["cadastros"] = f"Não foi possível ler os cadastros SIEG: {exc}"
    if ativos is not None and index is not None:
        out["certificados"] = [comparar(c, index.get(c["documento"], []) if not c["problema_documento"] else []) for c in ativos]
        order = {"sieg": 0, "pasta": 1, "sem_pasta": 2, "revisar": 3, "acompanhar": 4, "ok": 5}
        out["certificados"].sort(key=lambda x: (order[x["grupo"]], normalizar(x["nome"])))
    if ativos is not None:
        try:
            path, export_date = ultima_planilha(pasta_sieg, "Certidao-VisaoGeral")
            rows, info, avisos = carregar_certidoes(path, export_date, ativos)
            out["certidoes"] = rows
            out["fontes"]["certidoes"] = info
            out["avisos"].extend(avisos)
        except Exception as exc:
            out["erros"]["certidoes"] = f"Não foi possível ler as certidões SIEG: {exc}"
    certs = out["certificados"]
    pronto = ativos is not None and index is not None
    out["resumo"] = {"ativos": len(ativos) if ativos is not None else None,
                     **{g: sum(c["grupo"] == g for c in certs) if pronto else None for g in ("sieg", "pasta", "sem_pasta", "ok", "acompanhar", "revisar")},
                     "breve": sum(c["proximo_vencimento"] for c in certs) if pronto else None,
                     "estado": sum(not c["estado"] for c in ativos) if ativos is not None else None}
    out["certificados_disponiveis"] = pronto
    out["certidoes_disponiveis"] = "certidoes" in out["fontes"]
    for key in ("cadastros", "certidoes"):
        info = out["fontes"].get(key, {})
        if info.get("data_exportacao", "") > date.today().isoformat():
            out["avisos"].append(f"{info['arquivo']}: data de exportação futura. Confira o nome do arquivo.")
    return out
